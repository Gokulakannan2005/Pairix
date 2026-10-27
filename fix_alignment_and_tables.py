import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
import os
import docx2pdf

def format_cell_text(cell, font_size_pt=10, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(font_size_pt)
            if bold:
                r.font.bold = True

def apply_table_styles(table, col_widths_in):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_w = sum(col_widths_in)

    # Set tblW in tblPr
    tblPr = table._tbl.tblPr
    tblW = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblW')
    if tblW is not None:
        tblW.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'dxa')
        tblW.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', str(int(total_w * 1440)))

    # Set tblGrid
    tblGrid = table._tbl.tblGrid
    if tblGrid is not None:
        gridCols = tblGrid.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')
        for idx, w in enumerate(col_widths_in):
            if idx < len(gridCols):
                gridCols[idx].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', str(int(w * 1440)))

    # Repeat header row
    if len(table.rows) > 0:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        tblHeader = trPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader')
        if tblHeader is None:
            trPr.append(OxmlElement('w:tblHeader'))

    # Format rows
    for r_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        cantSplit = trPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit')
        if cantSplit is None:
            trPr.append(OxmlElement('w:cantSplit'))

        for c_idx, w in enumerate(col_widths_in):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell.width = Inches(w)
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW')
                if tcW is not None:
                    tcW.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'dxa')
                    tcW.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', str(int(w * 1440)))
                
                # Header formatting
                if r_idx == 0:
                    format_cell_text(cell, font_size_pt=10, bold=True)
                else:
                    format_cell_text(cell, font_size_pt=9.5)

def fix_report():
    doc = docx.Document('Pairix_Phase_I_Report_final.docx')
    print("Fixing table alignment and column widths...")

    # Table 0: Certificate signatures - DO NOT TOUCH!
    
    # Clean Table 9 if it still has 4 columns (remove empty col 2)
    t9 = doc.tables[9]
    if len(t9.columns) == 4:
        print("Removing empty column 2 from Table 9...")
        for row in t9.rows:
            row._tr.remove(row.cells[2]._tc)
        tblGrid = t9._tbl.tblGrid
        if tblGrid is not None:
            cols = tblGrid.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')
            if len(cols) > 2:
                tblGrid.remove(cols[2])

    # Table widths map (all sum to exactly 6.70 in, nicely centered within 6.90 in printable width)
    table_widths = {
        1: [1.3, 4.6, 0.8],           # TOC
        2: [1.2, 4.7, 0.8],           # List of Figures
        3: [1.2, 4.7, 0.8],           # List of Tables
        4: [1.8, 4.9],                # Abbreviations
        5: [1.7, 2.5, 2.5],           # Lit Review
        6: [1.4, 1.0, 2.3, 2.0],      # Profile Fields
        7: [1.8, 3.4, 1.5],           # Component Weights
        8: [2.0, 4.7],                # Auction Rules
        9: [2.4, 0.9, 3.4],           # API Endpoints (3 cols)
        10: [1.3, 3.4, 2.0],          # Implementation Plan
        11: [0.9, 2.4, 2.2, 1.2],     # Test Cases
        12: [1.8, 2.6, 2.3],          # Evaluation Metrics
        13: [1.3, 1.4, 1.0, 3.0],     # Data Dictionary (was 7.03 in, now 6.70 in!)
        14: [1.7, 1.0, 1.0, 1.0, 1.0, 1.0], # Sample Profiles (clean 1.0 in columns)
        15: [1.4, 2.6, 2.7],          # Edge Cases
        16: [1.2, 2.7, 2.8],          # Acceptance Priorities
    }

    for t_idx, widths in table_widths.items():
        if t_idx < len(doc.tables):
            apply_table_styles(doc.tables[t_idx], widths)
            print(f"Applied centered styles to Table {t_idx} (width={sum(widths):.2f} in)")

    # Ensure all image paragraphs and figure captions are centered
    print("Checking image and caption alignments...")
    for i, p in enumerate(doc.paragraphs):
        # If paragraph has drawing/image
        if any('drawing' in r._element.xml for r in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # If paragraph is a Figure caption
        if p.text.strip().startswith('Figure 4.'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10.5)

    output_docx = 'Pairix_Phase_I_Report_final.docx'
    print(f"Saving fixed Word document to {output_docx}...")
    doc.save(output_docx)
    print("Word document saved successfully!")

    # Now generate the PDF using docx2pdf
    output_pdf = 'Pairix_Phase_I_Report_final.pdf'
    print(f"Converting {output_docx} to {output_pdf}...")
    docx2pdf.convert(output_docx, output_pdf)
    print(f"PDF generated successfully! Size: {os.path.getsize(output_pdf)} bytes")

if __name__ == '__main__':
    fix_report()
